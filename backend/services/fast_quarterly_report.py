"""
Generate a quarterly FAST executive DOCX from FAST_Report + WAHIS INFUR.

Always targets the last completed calendar quarter (e.g. in September → Q2).
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from database import db_helper

# --- Diseases (display long names; DB stores short codes) ---
DISEASE_ORDER = ["FMD", "PPR", "LSD", "RVF", "SPGP"]
DISEASE_LONG = {
    "FMD": "Foot-and-Mouth Disease (FMD)",
    "PPR": "Peste des Petits Ruminants (PPR)",
    "LSD": "Lumpy Skin Disease (LSD)",
    "RVF": "Rift Valley Fever (RVF)",
    "SPGP": "Sheep Pox and Goat Pox (SPGP)",
}
DISEASE_RGB = {
    "FMD": RGBColor(0x1D, 0x4E, 0xD8),  # blue
    "PPR": RGBColor(0xEA, 0x58, 0x0C),  # orange
    "LSD": RGBColor(0x7C, 0x3A, 0xED),  # purple
    "RVF": RGBColor(0xDC, 0x26, 0x26),  # red
    "SPGP": RGBColor(0x16, 0xA3, 0x4A),  # green
}

REGIONS = [
    "Near East",
    "North Africa",
    "South East European Neighbourhood (SEEN)",
    "Europe",
]
REGION_SHORT = {
    "Near East": "Near East",
    "North Africa": "North Africa",
    "South East European Neighbourhood (SEEN)": "SEEN",
    "Europe": "Europe (INFUR)",
}

# EU / high-attention neighbourhood ( overlaps SEEN + known EU borders )
EU_ATTENTION_COUNTRIES = {
    "Türkiye",
    "Turkey",
    "Georgia",
    "Armenia",
    "Azerbaijan",
    "Ukraine",
    "Moldova",
    "Republic of Moldova",
    "Serbia",
    "Bosnia and Herzegovina",
    "Albania",
    "North Macedonia",
    "Montenegro",
    "Kosovo",
    "Belarus",
    "Russian Federation",
    "Russia",
}

_INFUR_DISEASE_MAP = (
    ("foot and mouth", "FMD"),
    ("lumpy skin", "LSD"),
    ("peste des petits", "PPR"),
    ("rift valley", "RVF"),
    ("sheep pox", "SPGP"),
    ("goat pox", "SPGP"),
)

_INFUR_NON_EUROPE = {
    "israel",
    "mauritania",
    "algeria",
    "lebanon",
    "palestine",
    "mali",
    "sudan",
    "south sudan (rep. of)",
    "south sudan",
}


def last_completed_quarter(today: Optional[date] = None) -> Tuple[int, int]:
    """Calendar rule: Apr–Jun→Q1, Jul–Sep→Q2, Oct–Dec→Q3, Jan–Mar→prev Q4."""
    d = today or date.today()
    m = d.month
    if m <= 3:
        return d.year - 1, 4
    if m <= 6:
        return d.year, 1
    if m <= 9:
        return d.year, 2
    return d.year, 3


def previous_quarter(year: int, quarter: int) -> Tuple[int, int]:
    if quarter == 1:
        return year - 1, 4
    return year, quarter - 1


def quarter_window(year: int, quarter: int, n: int = 4) -> List[Tuple[int, int]]:
    """Oldest → newest, ending at (year, quarter), length n."""
    out = [(year, quarter)]
    y, q = year, quarter
    for _ in range(n - 1):
        y, q = previous_quarter(y, q)
        out.append((y, q))
    out.reverse()
    return out


def _parse_outbreaks(val: Any) -> int:
    if val is None or val == "":
        return 0
    try:
        return max(0, int(float(str(val).strip())))
    except (TypeError, ValueError):
        return 0


def _norm_disease(raw: Any) -> Optional[str]:
    if not raw:
        return None
    s = str(raw).strip().upper()
    if s in DISEASE_LONG:
        return s
    for code in DISEASE_ORDER:
        if code in s or DISEASE_LONG[code].upper() in s:
            return code
    # long name without code
    lower = str(raw).lower()
    if "foot" in lower and "mouth" in lower:
        return "FMD"
    if "peste" in lower or "ppr" in lower:
        return "PPR"
    if "lumpy" in lower:
        return "LSD"
    if "rift" in lower:
        return "RVF"
    if "sheep pox" in lower or "goat pox" in lower or "spgp" in lower or "sgp" in lower:
        return "SPGP"
    return None


def _map_infur_disease(raw: Optional[str]) -> Optional[str]:
    name = (raw or "").lower()
    for needle, code in _INFUR_DISEASE_MAP:
        if needle in name:
            return code
    return None


def _traffic_light(outbreaks: int, countries: int) -> str:
    """Red / Orange / Green based on outbreak count and country spread."""
    if outbreaks <= 0:
        return "Green"
    if countries >= 3 or outbreaks >= 10:
        return "Red"
    return "Orange"


def _traffic_emoji(level: str) -> str:
    return {"Red": "🔴", "Orange": "🟠", "Green": "🟢"}.get(level, "⚪")


def _trend_arrow(current: int, previous: int) -> str:
    if previous <= 0:
        if current > 0:
            return "↑"
        return "→"
    change = (current - previous) / previous * 100
    if change > 5:
        return "↑"
    if change < -5:
        return "↓"
    return "→"


def _pct_change(current: int, previous: int) -> Optional[float]:
    if previous <= 0:
        return None if current == 0 else 100.0
    return (current - previous) / previous * 100.0


def _sparkline(values: List[int]) -> str:
    """Simple unicode bar sparkline for DOCX text."""
    if not values:
        return "—"
    mx = max(values) or 1
    blocks = "▁▂▃▄▅▆▇█"
    return "".join(blocks[min(7, int(v / mx * 7))] for v in values)


def _set_cell_shading(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


@dataclass
class DiseaseStats:
    code: str
    current: int = 0
    previous: int = 0
    countries: int = 0
    prev_countries: int = 0
    by_region: Dict[str, int] = field(default_factory=dict)
    history: List[int] = field(default_factory=list)  # last 4 quarters
    top_country: str = "—"
    vaccination_mentions: int = 0
    vaccination_doses: int = 0


@dataclass
class CountryDiseaseRow:
    country: str
    region: str
    disease: str
    outbreaks: int
    prev_outbreaks: int
    vaccination: str
    surveillance: str
    is_eu_attention: bool = False


async def _load_fast_rows(year: int, quarter: int) -> List[dict]:
    result = await db_helper.execute_main_query(
        """
        SELECT Year, Quarter, Region, Country, Disease,
               Outbreaks, Cases, Outbreak_Description, Epidemiological_Information,
               Surveillance, Vaccination, Vaccination_Doses, Vaccination_Description,
               Other_Info, Source
        FROM FAST_Report
        WHERE Year = %s AND Quarter = %s
          AND (Disease IS NULL OR Disease <> 'BEF')
        """,
        (year, quarter),
    )
    if result["error"]:
        raise RuntimeError(result["error"])
    return result["data"] or []


async def _load_infur_for_quarter(year: int, quarter: int) -> List[dict]:
    """INFUR Europe points whose outbreak start falls in the quarter."""
    start_m = {1: "01", 2: "04", 3: "07", 4: "10"}[quarter]
    end_m = {1: "03", 2: "06", 3: "09", 4: "12"}[quarter]
    end_d = {1: "31", 2: "30", 3: "30", 4: "31"}[quarter]
    start = f"{year}-{start_m}-01"
    end = f"{year}-{end_m}-{end_d}"

    result = await db_helper.execute_main_query(
        """
        SELECT country, disease, outbreakStartDate, eventStatus,
               cases, deaths, location, reason
        FROM INFUR
        WHERE outbreakStartDate IS NOT NULL
          AND outbreakStartDate >= %s
          AND outbreakStartDate <= %s
        ORDER BY outbreakStartDate DESC
        """,
        (start, end),
    )
    if result["error"]:
        raise RuntimeError(result["error"])

    points = []
    for row in result["data"] or []:
        country = (row.get("country") or "").strip()
        if not country or country.lower() in _INFUR_NON_EUROPE:
            continue
        code = _map_infur_disease(row.get("disease"))
        if not code:
            continue
        points.append(
            {
                "country": country,
                "disease": code,
                "outbreakStartDate": str(row.get("outbreakStartDate") or "")[:10],
                "cases": row.get("cases"),
                "location": row.get("location"),
                "reason": row.get("reason"),
            }
        )
    return points


def _aggregate(
    current_rows: List[dict],
    previous_rows: List[dict],
    history_by_period: Dict[Tuple[int, int], List[dict]],
    window: List[Tuple[int, int]],
    infur_current: List[dict],
    infur_previous: List[dict],
) -> Tuple[Dict[str, DiseaseStats], List[CountryDiseaseRow]]:
    stats: Dict[str, DiseaseStats] = {c: DiseaseStats(code=c) for c in DISEASE_ORDER}
    country_rows: List[CountryDiseaseRow] = []

    def accumulate(rows: List[dict], *, current: bool) -> None:
        by_dc: Dict[Tuple[str, str], dict] = {}
        for row in rows:
            code = _norm_disease(row.get("Disease"))
            if not code:
                continue
            country = str(row.get("Country") or "").strip() or "Unknown"
            region = str(row.get("Region") or "").strip() or "Other"
            key = (code, country)
            if key not in by_dc:
                by_dc[key] = {
                    "outbreaks": 0,
                    "region": region,
                    "vaccination": row.get("Vaccination"),
                    "vacc_desc": str(row.get("Vaccination_Description") or ""),
                    "vacc_doses": _parse_outbreaks(row.get("Vaccination_Doses")),
                    "surveillance": str(row.get("Surveillance") or "").strip(),
                }
            by_dc[key]["outbreaks"] += _parse_outbreaks(row.get("Outbreaks"))
            # Prefer non-empty surveillance / vacc text
            if row.get("Surveillance"):
                by_dc[key]["surveillance"] = str(row.get("Surveillance")).strip()
            if row.get("Vaccination_Description"):
                by_dc[key]["vacc_desc"] = str(row.get("Vaccination_Description")).strip()

        for (code, country), agg in by_dc.items():
            st = stats[code]
            if current:
                st.current += agg["outbreaks"]
                st.by_region[agg["region"]] = st.by_region.get(agg["region"], 0) + agg["outbreaks"]
                if agg["outbreaks"] > 0:
                    st.countries += 1
                if agg["vacc_doses"] > 0 or str(agg.get("vaccination") or "") in ("1", "Yes", "yes", "true"):
                    st.vaccination_mentions += 1
                    st.vaccination_doses += agg["vacc_doses"]
                vacc_status = agg["vacc_desc"] or (
                    f"Doses: {agg['vacc_doses']}" if agg["vacc_doses"] else str(agg.get("vaccination") or "—")
                )
                country_rows.append(
                    CountryDiseaseRow(
                        country=country,
                        region=agg["region"],
                        disease=code,
                        outbreaks=agg["outbreaks"],
                        prev_outbreaks=0,
                        vaccination=vacc_status[:120],
                        surveillance=agg["surveillance"][:120] or "—",
                        is_eu_attention=country in EU_ATTENTION_COUNTRIES,
                    )
                )
            else:
                st.previous += agg["outbreaks"]
                if agg["outbreaks"] > 0:
                    st.prev_countries += 1

    accumulate(current_rows, current=True)
    accumulate(previous_rows, current=False)

    # Attach previous outbreak counts onto country rows
    prev_map: Dict[Tuple[str, str], int] = defaultdict(int)
    for row in previous_rows:
        code = _norm_disease(row.get("Disease"))
        if not code:
            continue
        country = str(row.get("Country") or "").strip() or "Unknown"
        prev_map[(code, country)] += _parse_outbreaks(row.get("Outbreaks"))
    for cr in country_rows:
        cr.prev_outbreaks = prev_map.get((cr.disease, cr.country), 0)

    # INFUR Europe contribution for current + previous
    for pt in infur_current:
        code = pt["disease"]
        if code not in stats:
            continue
        stats[code].current += 1
        stats[code].by_region["Europe"] = stats[code].by_region.get("Europe", 0) + 1
        stats[code].countries += 1  # may over-count countries; refine below
    # Recompute Europe country uniqueness for current INFUR
    for code in DISEASE_ORDER:
        eu_countries = {p["country"] for p in infur_current if p["disease"] == code}
        # countries already counted FAST countries; add EU-only extras roughly
        # Simpler: set countries as unique FAST countries with outbreaks + INFUR countries
        pass

    # Recalculate countries properly
    for code in DISEASE_ORDER:
        fast_countries = {
            cr.country for cr in country_rows if cr.disease == code and cr.outbreaks > 0
        }
        infur_countries = {p["country"] for p in infur_current if p["disease"] == code}
        stats[code].countries = len(fast_countries | infur_countries)
        # Fix current count: FAST sum + INFUR event count
        fast_sum = sum(cr.outbreaks for cr in country_rows if cr.disease == code)
        infur_sum = sum(1 for p in infur_current if p["disease"] == code)
        stats[code].current = fast_sum + infur_sum
        prev_fast = stats[code].previous
        prev_infur = sum(1 for p in infur_previous if p["disease"] == code)
        stats[code].previous = prev_fast + prev_infur

    # Add INFUR country rows for annex / snapshots
    infur_by_dc: Dict[Tuple[str, str], int] = defaultdict(int)
    for p in infur_current:
        infur_by_dc[(p["disease"], p["country"])] += 1
    existing = {(cr.disease, cr.country) for cr in country_rows}
    for (code, country), n in infur_by_dc.items():
        if (code, country) in existing:
            # merge into existing if same name
            for cr in country_rows:
                if cr.disease == code and cr.country == country:
                    cr.outbreaks += n
                    break
        else:
            country_rows.append(
                CountryDiseaseRow(
                    country=country,
                    region="Europe",
                    disease=code,
                    outbreaks=n,
                    prev_outbreaks=sum(
                        1 for p in infur_previous if p["disease"] == code and p["country"] == country
                    ),
                    vaccination="— (INFUR)",
                    surveillance="WAHIS INFUR",
                    is_eu_attention=True,
                )
            )

    # History for sparklines
    for code in DISEASE_ORDER:
        hist = []
        for y, q in window:
            rows = history_by_period.get((y, q), [])
            total = sum(
                _parse_outbreaks(r.get("Outbreaks"))
                for r in rows
                if _norm_disease(r.get("Disease")) == code
            )
            # INFUR only added for current/previous windows ends — approx from loaded periods
            hist.append(total)
        # Replace last two with fully adjusted current/previous if window ends there
        if len(hist) >= 1:
            hist[-1] = stats[code].current
        if len(hist) >= 2:
            hist[-2] = stats[code].previous
        stats[code].history = hist

        # Top country
        tops = sorted(
            [cr for cr in country_rows if cr.disease == code and cr.outbreaks > 0],
            key=lambda x: x.outbreaks,
            reverse=True,
        )
        stats[code].top_country = tops[0].country if tops else "—"

    return stats, country_rows


def _key_messages(
    stats: Dict[str, DiseaseStats],
    country_rows: List[CountryDiseaseRow],
    year: int,
    quarter: int,
) -> List[str]:
    msgs: List[Tuple[int, str]] = []

    for code, st in stats.items():
        arrow = _trend_arrow(st.current, st.previous)
        if st.current > 0 and arrow == "↑":
            pct = _pct_change(st.current, st.previous)
            pct_txt = f" ({pct:.0f}% vs previous quarter)" if pct is not None else ""
            msgs.append(
                (
                    st.current + 50,
                    f"{DISEASE_LONG[code]}: {st.current} outbreak reports this quarter{pct_txt}; "
                    f"highest activity in {st.top_country}.",
                )
            )
        elif st.current >= 10:
            msgs.append(
                (
                    st.current,
                    f"{DISEASE_LONG[code]}: elevated activity ({st.current} reports across "
                    f"{st.countries} countries); focus on {st.top_country}.",
                )
            )

    # EU-attention spikes
    for cr in country_rows:
        if not cr.is_eu_attention or cr.outbreaks <= 0:
            continue
        if cr.prev_outbreaks > 0 and cr.outbreaks > cr.prev_outbreaks * 1.05:
            msgs.append(
                (
                    40 + cr.outbreaks,
                    f"Cross-border attention — {cr.country} ({DISEASE_LONG[cr.disease]}): "
                    f"{cr.outbreaks} reports (up from {cr.prev_outbreaks}).",
                )
            )
        elif cr.outbreaks >= 3:
            msgs.append(
                (
                    30 + cr.outbreaks,
                    f"EU-neighbourhood concern — {cr.country}: {cr.outbreaks} "
                    f"{DISEASE_LONG[cr.disease]} reports this quarter.",
                )
            )

    # Vaccination gaps in high-outbreak countries
    for cr in country_rows:
        if cr.outbreaks < 3:
            continue
        vacc = (cr.vaccination or "").lower()
        if vacc in ("—", "0", "no", "none", "") or "0%" in vacc:
            msgs.append(
                (
                    20 + cr.outbreaks,
                    f"Control gap — {cr.country} reports {cr.outbreaks} "
                    f"{DISEASE_LONG[cr.disease]} outbreaks with limited/unclear vaccination information.",
                )
            )

    # Stable / improving overall note if few spikes
    total_cur = sum(s.current for s in stats.values())
    total_prev = sum(s.previous for s in stats.values())
    if total_cur == 0:
        msgs.append((5, f"No FAST outbreak reports recorded for {year}-Q{quarter} in the FAST table / INFUR filter."))
    elif _trend_arrow(total_cur, total_prev) == "↓":
        msgs.append(
            (
                10,
                f"Overall FAST outbreak reports declined versus the previous quarter "
                f"({total_cur} vs {total_prev}).",
            )
        )

    msgs.sort(key=lambda x: x[0], reverse=True)
    # Dedupe similar starts
    seen = set()
    out = []
    for _, text in msgs:
        key = text[:48]
        if key in seen:
            continue
        seen.add(key)
        out.append(text)
        if len(out) >= 5:
            break
    while len(out) < 5:
        out.append("No additional critical signals flagged by automated rules for this quarter.")
    return out[:5]


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _add_para(doc: Document, text: str, *, bold: bool = False, size: int = 11) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def _page_break(doc: Document) -> None:
    doc.add_page_break()


def build_docx(
    *,
    year: int,
    quarter: int,
    stats: Dict[str, DiseaseStats],
    country_rows: List[CountryDiseaseRow],
    key_messages: List[str],
) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    prev_y, prev_q = previous_quarter(year, quarter)

    # ----- Page 1: Executive Dashboard -----
    _add_heading(doc, f"FAST Quarterly Executive Report — {year} Q{quarter}", 0)
    _add_para(
        doc,
        f"Last completed calendar quarter. Trends vs {prev_y} Q{prev_q}. "
        "Sources: FAST_Report (neighbourhood) and WAHIS INFUR (Europe).",
        size=10,
    )
    _add_heading(doc, "1. Executive Dashboard (30-second read)", 1)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    headers = ["Disease", "Outbreaks this quarter", "Trend", "Highest risk area", "Status"]
    for i, h in enumerate(headers):
        hdr[i].text = h
        _set_cell_shading(hdr[i], "1F4E79")
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.bold = True
                r.font.size = Pt(9)

    for code in DISEASE_ORDER:
        st = stats[code]
        level = _traffic_light(st.current, st.countries)
        row = table.add_row().cells
        row[0].text = DISEASE_LONG[code]
        for p in row[0].paragraphs:
            for r in p.runs:
                r.font.color.rgb = DISEASE_RGB[code]
                r.bold = True
                r.font.size = Pt(9)
        row[1].text = str(st.current)
        arrow = _trend_arrow(st.current, st.previous)
        pct = _pct_change(st.current, st.previous)
        pct_s = f" {pct:+.0f}%" if pct is not None else ""
        row[2].text = f"{arrow}{pct_s} (prev {st.previous})"
        row[3].text = st.top_country
        row[4].text = f"{_traffic_emoji(level)} {level}"
        fill = {"Red": "FECACA", "Orange": "FED7AA", "Green": "BBF7D0"}[level]
        _set_cell_shading(row[4], fill)

    _add_para(
        doc,
        "Traffic lights: 🔴 Active spread (≥3 countries or ≥10 reports) · "
        "🟠 Localized (outbreaks present but limited) · 🟢 No outbreaks / monitoring.",
        size=9,
    )

    # ----- Page 2: Key Messages -----
    _page_break(doc)
    _add_heading(doc, "2. Key Messages", 1)
    _add_para(doc, "Five headlines for focal points and risk managers:", size=10)
    for i, msg in enumerate(key_messages, 1):
        _add_para(doc, f"{i}. {msg}", size=11)

    # ----- Page 3: Regional heat map placeholder -----
    _page_break(doc)
    _add_heading(doc, "3. Regional Risk Overview", 1)
    _add_para(
        doc,
        "[MAP PLACEHOLDER] Insert regional risk heat map here (copy/paste from GIS or Fast Report map). "
        "Highlight EU-border / hub countries as needed.",
        size=10,
    )
    _add_para(doc, "Regional outbreak totals by disease (FAST regions + Europe INFUR):", bold=True)

    rtable = doc.add_table(rows=1, cols=1 + len(DISEASE_ORDER))
    rtable.style = "Table Grid"
    rh = rtable.rows[0].cells
    rh[0].text = "Region"
    for i, code in enumerate(DISEASE_ORDER, 1):
        rh[i].text = code
        for p in rh[i].paragraphs:
            for r in p.runs:
                r.font.color.rgb = DISEASE_RGB[code]
                r.bold = True
    for region in REGIONS:
        row = rtable.add_row().cells
        row[0].text = REGION_SHORT.get(region, region)
        for i, code in enumerate(DISEASE_ORDER, 1):
            n = stats[code].by_region.get(region, 0)
            # Europe stored as "Europe"
            if region == "Europe":
                n = stats[code].by_region.get("Europe", 0)
            row[i].text = str(n)
            if n >= 10:
                _set_cell_shading(row[i], "FECACA")
            elif n > 0:
                _set_cell_shading(row[i], "FED7AA")
            else:
                _set_cell_shading(row[i], "BBF7D0")

    _add_para(doc, "EU-attention countries with outbreaks this quarter:", bold=True, size=10)
    eu_hit = [cr for cr in country_rows if cr.is_eu_attention and cr.outbreaks > 0]
    if not eu_hit:
        _add_para(doc, "None flagged in automated EU-attention list.", size=10)
    else:
        for cr in sorted(eu_hit, key=lambda x: x.outbreaks, reverse=True)[:15]:
            _add_para(
                doc,
                f"• {cr.country} — {DISEASE_LONG[cr.disease]}: {cr.outbreaks} "
                f"(trend {_trend_arrow(cr.outbreaks, cr.prev_outbreaks)})",
                size=10,
            )

    # ----- Pages 4–8: Disease snapshots -----
    for code in DISEASE_ORDER:
        st = stats[code]
        _page_break(doc)
        _add_heading(doc, f"Disease Snapshot — {DISEASE_LONG[code]}", 1)
        level = _traffic_light(st.current, st.countries)
        p = doc.add_paragraph()
        run = p.add_run(f"Accent colour · Status {_traffic_emoji(level)} {level}")
        run.font.color.rgb = DISEASE_RGB[code]
        run.bold = True

        _add_para(
            doc,
            f"Outbreaks this quarter: {st.current}  |  Previous: {st.previous}  |  "
            f"Trend: {_trend_arrow(st.current, st.previous)}  |  Countries: {st.countries}  |  "
            f"Top area: {st.top_country}",
            size=10,
        )
        _add_para(
            doc,
            f"4-quarter sparkline (oldest→newest): {_sparkline(st.history)}  "
            f"Values: {', '.join(str(v) for v in st.history)}",
            size=10,
        )
        _add_para(
            doc,
            f"Vaccination rows with activity/doses (FAST): {st.vaccination_mentions} countries/rows; "
            f"sum of reported doses (where numeric): {st.vaccination_doses}.",
            size=9,
        )

        _add_para(doc, "Condensed country matrix (noise filtered):", bold=True, size=10)
        # Include if outbreaks>0 OR eu attention OR previous outbreaks
        snapshot = [
            cr
            for cr in country_rows
            if cr.disease == code
            and (cr.outbreaks > 0 or cr.is_eu_attention or cr.prev_outbreaks > 0)
        ]
        snapshot.sort(key=lambda x: (0 if x.outbreaks > 0 else 1, -x.outbreaks, x.country))

        ctable = doc.add_table(rows=1, cols=5)
        ctable.style = "Table Grid"
        ch = ctable.rows[0].cells
        for i, h in enumerate(["Country", "Outbreaks ☣", "Vaccination 💉", "Surveillance 🔍", "Trend"]):
            ch[i].text = h
            _set_cell_shading(ch[i], "E5E7EB")
        if not snapshot:
            row = ctable.add_row().cells
            row[0].text = "No countries meeting inclusion rules"
            row[1].text = "0"
            row[2].text = "—"
            row[3].text = "—"
            row[4].text = "→"
        else:
            for cr in snapshot[:25]:
                row = ctable.add_row().cells
                row[0].text = cr.country + (" *" if cr.is_eu_attention else "")
                row[1].text = str(cr.outbreaks)
                row[2].text = cr.vaccination[:80]
                row[3].text = cr.surveillance[:80]
                row[4].text = _trend_arrow(cr.outbreaks, cr.prev_outbreaks)
        _add_para(doc, "* = EU-attention / neighbourhood priority country.", size=8)

    # ----- Annexes: full country tables -----
    _page_break(doc)
    _add_heading(doc, "Country Detail Annexes", 1)
    _add_para(
        doc,
        "Full technical tables for veterinary officers (all FAST + INFUR-linked rows for the quarter).",
        size=10,
    )

    for code in DISEASE_ORDER:
        _add_heading(doc, DISEASE_LONG[code], 2)
        annex = [cr for cr in country_rows if cr.disease == code]
        annex.sort(key=lambda x: (-x.outbreaks, x.country))
        at = doc.add_table(rows=1, cols=6)
        at.style = "Table Grid"
        ah = at.rows[0].cells
        for i, h in enumerate(
            ["Country", "Region", "Outbreaks", "Prev Q", "Vaccination", "Surveillance"]
        ):
            ah[i].text = h
            _set_cell_shading(ah[i], "DBEAFE")
        if not annex:
            r = at.add_row().cells
            r[0].text = "—"
            r[2].text = "0"
        else:
            for cr in annex:
                r = at.add_row().cells
                r[0].text = cr.country
                r[1].text = REGION_SHORT.get(cr.region, cr.region)
                r[2].text = str(cr.outbreaks)
                r[3].text = str(cr.prev_outbreaks)
                r[4].text = cr.vaccination[:100]
                r[5].text = cr.surveillance[:100]
        doc.add_paragraph()

    # ----- Emerging issues & methodology -----
    _page_break(doc)
    _add_heading(doc, "Emerging Issues & Methodology", 1)
    _add_heading(doc, "Emerging issues outside / at the edge of the FAST neighbourhood", 2)
    _add_para(
        doc,
        "Europe (WAHIS INFUR) events in this quarter are summarised above under the Europe region. "
        "Review INFUR for EU Member States and near-border events (e.g. Mediterranean, Caucasus) "
        "and paste any additional narrative or map callouts here after generation.",
        size=10,
    )
    eu_infur = [cr for cr in country_rows if cr.region == "Europe" and cr.outbreaks > 0]
    if eu_infur:
        _add_para(doc, "INFUR Europe highlights this quarter:", bold=True, size=10)
        for cr in sorted(eu_infur, key=lambda x: -x.outbreaks)[:12]:
            _add_para(doc, f"• {cr.country} — {DISEASE_LONG[cr.disease]}: {cr.outbreaks}", size=10)

    _add_heading(doc, "Methodology (brief)", 2)
    _add_para(
        doc,
        "• Scope: last completed calendar quarter only (report period fixed at generation time).\n"
        "• FAST data: FAST_Report rows for the selected Year/Quarter (BEF excluded).\n"
        "• Europe: WAHIS INFUR rows with outbreak start date inside the quarter; non-European "
        "countries filtered out.\n"
        "• Trends: compare outbreak totals to the immediately previous calendar quarter; "
        "↑ if >5% increase, ↓ if >5% decrease, else →.\n"
        "• Traffic lights: Green = 0 outbreaks; Orange = outbreaks with <3 countries and <10 reports; "
        "Red = ≥3 countries or ≥10 reports.\n"
        "• Snapshot filtering: countries with 0 outbreaks are omitted unless EU-attention or "
        "they had outbreaks in the previous quarter.\n"
        "• Sparklines: last four calendar quarters ending at the report quarter "
        "(FAST totals; current/previous adjusted with INFUR).",
        size=9,
    )
    _add_para(
        doc,
        "This document is generated for review. Adapt narrative, paste maps, then export to PDF.",
        size=9,
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def generate_quarterly_fast_docx(today: Optional[date] = None) -> Tuple[bytes, int, int]:
    year, quarter = last_completed_quarter(today)
    prev_y, prev_q = previous_quarter(year, quarter)
    window = quarter_window(year, quarter, 4)

    history_by_period: Dict[Tuple[int, int], List[dict]] = {}
    for y, q in window:
        history_by_period[(y, q)] = await _load_fast_rows(y, q)

    current_rows = history_by_period[(year, quarter)]
    previous_rows = history_by_period.get((prev_y, prev_q), await _load_fast_rows(prev_y, prev_q))

    infur_current = await _load_infur_for_quarter(year, quarter)
    infur_previous = await _load_infur_for_quarter(prev_y, prev_q)

    stats, country_rows = _aggregate(
        current_rows,
        previous_rows,
        history_by_period,
        window,
        infur_current,
        infur_previous,
    )
    messages = _key_messages(stats, country_rows, year, quarter)
    docx_bytes = build_docx(
        year=year,
        quarter=quarter,
        stats=stats,
        country_rows=country_rows,
        key_messages=messages,
    )
    return docx_bytes, year, quarter
