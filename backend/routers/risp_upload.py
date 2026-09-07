"""Parse and import RISP/SOI bulk Excel uploads."""

from __future__ import annotations

import json
from datetime import date, datetime
from io import BytesIO
from typing import Any, Callable, Dict, List, Optional, Tuple

import openpyxl
from fastapi import HTTPException, UploadFile

from routers import risp_templates

VACCINATION_TYPE_BY_LABEL = {
    label: code for code, label in risp_templates.VACCINATION_TYPES
}
VACCINATION_TYPE_BY_LABEL.update({code: code for code, _ in risp_templates.VACCINATION_TYPES})

SHEET_NAMES = {
    "outbreaks": ("Outbreaks",),
    "vaccination": ("Vaccination",),
    "marketprice": ("Market prices", "Marketprice"),
}


def _cell_str(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell_int(value: Any) -> Optional[int]:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _cell_float(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _cell_date(value: Any) -> Optional[str]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text


def _split_list_field(value: Any) -> List[str]:
    text = _cell_str(value)
    if not text:
        return []
    if ";" in text:
        parts = [p.strip() for p in text.split(";")]
    elif "," in text:
        parts = [p.strip() for p in text.split(",")]
    else:
        parts = [text]
    return [p for p in parts if p]


def _read_headers(ws) -> Dict[str, int]:
    headers: Dict[str, int] = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        if cell.value:
            headers[str(cell.value).strip().lower()] = col_idx
    return headers


def _get_cell(ws, row: int, headers: Dict[str, int], key: str) -> Any:
    col = headers.get(key)
    if not col:
        return None
    return ws.cell(row=row, column=col).value


def _pick_sheet(wb, category: str):
    for name in SHEET_NAMES.get(category, ()):
        if name in wb.sheetnames:
            return wb[name]
    return wb.active


async def _read_workbook(file: UploadFile):
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file provided")
    lower = file.filename.lower()
    if not (lower.endswith(".xlsx") or lower.endswith(".xls")):
        raise HTTPException(status_code=400, detail="Only Excel files (.xlsx) are allowed")
    contents = await file.read()
    return openpyxl.load_workbook(BytesIO(contents), data_only=True)


def _resolve_district(
    cursor,
    country_id: int,
    district_id: Optional[int],
    district_name: Optional[str],
    province_id: Optional[int],
) -> Tuple[Optional[int], Optional[int], Optional[str]]:
    if district_id:
        cursor.execute(
            """
            SELECT d.id AS district_id, d.name AS district_name, p.id AS province_id
            FROM districts d
            JOIN provinces p ON p.id = d.province_id
            WHERE d.id = %s AND p.country_id = %s
            LIMIT 1
            """,
            (district_id, country_id),
        )
        row = cursor.fetchone()
        if row:
            return row["district_id"], row["province_id"], row["district_name"]
        return district_id, province_id, district_name

    if not district_name:
        return None, province_id, None

    params: list = [country_id, district_name]
    query = """
        SELECT d.id AS district_id, d.name AS district_name, p.id AS province_id
        FROM districts d
        JOIN provinces p ON p.id = d.province_id
        WHERE p.country_id = %s AND d.name = %s
    """
    if province_id:
        query += " AND p.id = %s"
        params.append(province_id)
    query += " ORDER BY d.id LIMIT 1"
    cursor.execute(query, tuple(params))
    row = cursor.fetchone()
    if row:
        return row["district_id"], row["province_id"], row["district_name"]
    return None, province_id, district_name


def _import_outbreaks(
    ws,
    *,
    cursor,
    user_id: int,
    country: str,
    program: str,
    country_id: Optional[int],
    is_soi: bool,
    default_year: Optional[str],
    default_quarter: Optional[str],
) -> Tuple[int, List[Dict[str, Any]]]:
    headers = _read_headers(ws)
    required = {"year", "quarter", "disease"}
    missing = sorted(required - set(headers))
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing required columns: {', '.join(missing)}",
        )

    imported = 0
    errors: List[Dict[str, Any]] = []
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for row_idx in range(2, ws.max_row + 1):
        disease = _cell_str(_get_cell(ws, row_idx, headers, "disease"))
        if not disease:
            continue

        year = _cell_str(_get_cell(ws, row_idx, headers, "year")) or default_year
        quarter = _cell_str(_get_cell(ws, row_idx, headers, "quarter")) or default_quarter
        if not year or not quarter:
            errors.append({"row": row_idx, "error": "year and quarter are required"})
            continue

        number_outbreaks = _cell_int(_get_cell(ws, row_idx, headers, "number_outbreaks"))
        if number_outbreaks is None:
            number_outbreaks = 0

        location = _cell_str(_get_cell(ws, row_idx, headers, "location"))
        district_name = _cell_str(_get_cell(ws, row_idx, headers, "district"))
        province_id = _cell_int(_get_cell(ws, row_idx, headers, "province_id"))
        district_id = _cell_int(_get_cell(ws, row_idx, headers, "district_id"))

        if is_soi and country_id:
            resolved_district_id, resolved_province_id, resolved_name = _resolve_district(
                cursor, country_id, district_id, district_name or location, province_id
            )
            if not resolved_district_id:
                errors.append(
                    {
                        "row": row_idx,
                        "error": f"Could not match district '{district_name or location}' for your country",
                    }
                )
                continue
            district_id = resolved_district_id
            province_id = resolved_province_id
            location = resolved_name or district_name or location
        elif not location and district_name:
            location = district_name

        areas = [location] if location else []
        status = _split_list_field(_get_cell(ws, row_idx, headers, "status"))
        serotype = _split_list_field(_get_cell(ws, row_idx, headers, "serotype"))
        species = _split_list_field(_get_cell(ws, row_idx, headers, "species"))
        control_measures = _split_list_field(_get_cell(ws, row_idx, headers, "control_measures"))
        comments = _cell_str(_get_cell(ws, row_idx, headers, "comments")) or ""
        date_suspected = _cell_date(_get_cell(ws, row_idx, headers, "date_suspected"))
        date_confirmed = _cell_date(_get_cell(ws, row_idx, headers, "date_confirmed"))
        latitude = _cell_float(_get_cell(ws, row_idx, headers, "latitude"))
        longitude = _cell_float(_get_cell(ws, row_idx, headers, "longitude"))

        match_sql = """
            SELECT id FROM risp_outbreaks
            WHERE user_id = %s AND year = %s AND quarter = %s AND disease_name = %s
        """
        match_params: list = [user_id, year, quarter, disease]
        if district_id is not None:
            match_sql += " AND district_id = %s"
            match_params.append(district_id)
        elif location:
            match_sql += " AND location = %s"
            match_params.append(location)
        else:
            match_sql += " AND (district_id IS NULL OR district_id = 0) AND (location IS NULL OR location = '')"
        match_sql += " ORDER BY id LIMIT 1"
        cursor.execute(match_sql, tuple(match_params))
        existing = cursor.fetchone()

        data_values = (
            number_outbreaks,
            json.dumps(areas),
            location,
            json.dumps(status),
            json.dumps(serotype),
            json.dumps(species),
            json.dumps(control_measures),
            comments,
            date_suspected,
            date_confirmed,
            latitude,
            longitude,
            province_id,
            district_id,
            program,
            "public",
        )

        if existing:
            row_id = existing["id"] if isinstance(existing, dict) else existing[0]
            cursor.execute(
                """
                UPDATE risp_outbreaks SET
                  number_outbreaks = %s, locations = %s, location = %s,
                  status = %s, serotype = %s, species = %s, control_measures = %s,
                  additional_info = %s, date_suspected = %s, date_confirmed = %s,
                  latitude = %s, longitude = %s, province_id = %s, district_id = %s,
                  program = %s, visibility = %s,
                  created_at = COALESCE(created_at, %s),
                  updated_at = %s
                WHERE id = %s AND user_id = %s
                """,
                data_values + (now, now, row_id, user_id),
            )
        else:
            cursor.execute(
                """
                INSERT INTO risp_outbreaks
                  (user_id, country, year, quarter, disease_name, number_outbreaks,
                   locations, location, status, serotype, species, control_measures, additional_info,
                   date_suspected, date_confirmed, latitude, longitude,
                   province_id, district_id, program, visibility, created_at, updated_at)
                VALUES
                  (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                   %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    user_id,
                    country,
                    year,
                    quarter,
                    disease,
                )
                + data_values
                + (now, now),
            )
        imported += 1

    return imported, errors


def _import_vaccinations(
    ws,
    *,
    cursor,
    user_id: int,
    country: str,
    program: str,
    country_id: Optional[int],
    is_soi: bool,
    default_year: Optional[str],
) -> Tuple[int, List[Dict[str, Any]]]:
    headers = _read_headers(ws)
    if "disease" not in headers and "disease_name" in headers:
        headers["disease"] = headers["disease_name"]
    if "disease" not in headers:
        raise HTTPException(status_code=400, detail="Missing required column: disease")

    imported = 0
    errors: List[Dict[str, Any]] = []
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for row_idx in range(2, ws.max_row + 1):
        disease = _cell_str(_get_cell(ws, row_idx, headers, "disease"))
        if not disease:
            continue

        year = _cell_str(_get_cell(ws, row_idx, headers, "year")) or default_year
        if not year:
            errors.append({"row": row_idx, "error": "year is required"})
            continue

        vaccination_type_label = _cell_str(_get_cell(ws, row_idx, headers, "vaccination_type"))
        vaccination_type = (
            VACCINATION_TYPE_BY_LABEL.get(vaccination_type_label or "", vaccination_type_label)
            if vaccination_type_label
            else None
        )

        location = _cell_str(_get_cell(ws, row_idx, headers, "location"))
        district_name = _cell_str(_get_cell(ws, row_idx, headers, "district"))
        province_id = _cell_int(_get_cell(ws, row_idx, headers, "province_id"))
        district_id = _cell_int(_get_cell(ws, row_idx, headers, "district_id"))

        if is_soi and country_id:
            resolved_district_id, resolved_province_id, resolved_name = _resolve_district(
                cursor, country_id, district_id, district_name or location, province_id
            )
            if not resolved_district_id:
                errors.append(
                    {
                        "row": row_idx,
                        "error": f"Could not match district '{district_name or location}' for your country",
                    }
                )
                continue
            district_id = resolved_district_id
            province_id = resolved_province_id
            location = resolved_name or district_name or location
        elif not location and district_name:
            location = district_name

        areas = [location] if location else []
        q1 = _cell_int(_get_cell(ws, row_idx, headers, "q1")) or 0
        q2 = _cell_int(_get_cell(ws, row_idx, headers, "q2")) or 0
        q3 = _cell_int(_get_cell(ws, row_idx, headers, "q3")) or 0
        q4 = _cell_int(_get_cell(ws, row_idx, headers, "q4")) or 0
        total = q1 + q2 + q3 + q4
        coverage = _cell_int(_get_cell(ws, row_idx, headers, "coverage")) or 0
        species = _split_list_field(_get_cell(ws, row_idx, headers, "species"))
        vaccine_details = _cell_str(_get_cell(ws, row_idx, headers, "vaccine_details"))
        status = _cell_str(_get_cell(ws, row_idx, headers, "status"))

        match_sql = """
            SELECT id FROM risp_vaccination
            WHERE user_id = %s AND year = %s AND disease_name = %s
              AND (location IS NULL OR location <> '__deleted__')
        """
        match_params: list = [user_id, year, disease]
        if district_id is not None:
            match_sql += " AND district_id = %s"
            match_params.append(district_id)
        elif location:
            match_sql += " AND location = %s"
            match_params.append(location)
        match_sql += " ORDER BY id LIMIT 1"
        cursor.execute(match_sql, tuple(match_params))
        existing = cursor.fetchone()

        values = (
            country,
            disease,
            year,
            status,
            vaccination_type,
            json.dumps(areas),
            location,
            province_id,
            district_id,
            program,
            "public",
            json.dumps(species),
            vaccine_details,
            q1,
            q2,
            q3,
            q4,
            total,
            coverage,
        )

        if existing:
            row_id = existing["id"] if isinstance(existing, dict) else existing[0]
            cursor.execute(
                """
                UPDATE risp_vaccination SET
                  country = %s, disease_name = %s, year = %s, status = %s,
                  vaccination_type = %s, geographical_areas = %s, location = %s,
                  province_id = %s, district_id = %s, program = %s, visibility = %s,
                  species = %s, vaccine_details = %s,
                  q1 = %s, q2 = %s, q3 = %s, q4 = %s, total = %s, coverage = %s,
                  created_at = COALESCE(created_at, %s)
                WHERE id = %s AND user_id = %s
                """,
                values + (now, row_id, user_id),
            )
        else:
            cursor.execute(
                """
                INSERT INTO risp_vaccination
                  (user_id, country, disease_name, year, status, vaccination_type,
                   geographical_areas, location, province_id, district_id, program, visibility,
                   species, vaccine_details, q1, q2, q3, q4, total, coverage, created_at)
                VALUES
                  (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (user_id,) + values + (now,),
            )
        imported += 1

    return imported, errors


def _import_market_prices(
    ws,
    *,
    cursor,
    user_id: int,
    country: str,
    program: str,
    country_id: Optional[int],
    is_soi: bool,
    default_year: Optional[str],
    default_quarter: Optional[str],
) -> Tuple[int, List[Dict[str, Any]]]:
    headers = _read_headers(ws)
    required = {"species", "market_level", "product"}
    missing = sorted(required - set(headers))
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing required columns: {', '.join(missing)}",
        )

    imported = 0
    errors: List[Dict[str, Any]] = []
    deleted_marker = "__deleted__"
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for row_idx in range(2, ws.max_row + 1):
        species = (_cell_str(_get_cell(ws, row_idx, headers, "species")) or "").lower()
        market_level = (_cell_str(_get_cell(ws, row_idx, headers, "market_level")) or "").lower()
        product = (_cell_str(_get_cell(ws, row_idx, headers, "product")) or "").lower()
        if not species and not market_level and not product:
            continue
        if species not in risp_templates.MARKET_SPECIES:
            errors.append({"row": row_idx, "error": f"Invalid species '{species}'"})
            continue
        if market_level not in risp_templates.MARKET_LEVELS:
            errors.append({"row": row_idx, "error": f"Invalid market_level '{market_level}'"})
            continue
        if product not in risp_templates.MARKET_PRODUCTS:
            errors.append({"row": row_idx, "error": f"Invalid product '{product}'"})
            continue

        year = _cell_str(_get_cell(ws, row_idx, headers, "year")) or default_year
        quarter = _cell_str(_get_cell(ws, row_idx, headers, "quarter")) or default_quarter
        if not year or not quarter:
            errors.append({"row": row_idx, "error": "year and quarter are required"})
            continue

        price_min = _cell_float(_get_cell(ws, row_idx, headers, "price_min"))
        price_max = _cell_float(_get_cell(ws, row_idx, headers, "price_max"))
        price_avg = _cell_float(_get_cell(ws, row_idx, headers, "price_avg"))
        if price_min is None and price_max is None and price_avg is None:
            continue

        reference = _cell_str(_get_cell(ws, row_idx, headers, "reference"))

        cursor.execute(
            """
            SELECT id FROM risp_marketprice
            WHERE user_id = %s AND year = %s AND quarter = %s
              AND species = %s AND market_level = %s AND product = %s
              AND (reference IS NULL OR reference <> %s)
            ORDER BY id LIMIT 1
            """,
            (
                user_id,
                year,
                quarter,
                species,
                market_level,
                product,
                deleted_marker,
            ),
        )
        existing = cursor.fetchone()
        values = (
            species,
            market_level,
            product,
            price_min,
            price_max,
            price_avg,
            reference,
            program,
        )

        if existing:
            row_id = existing["id"] if isinstance(existing, dict) else existing[0]
            cursor.execute(
                """
                UPDATE risp_marketprice SET
                  species = %s, market_level = %s, product = %s,
                  price_min = %s, price_max = %s, price_avg = %s,
                  reference = %s, program = %s,
                  visibility = 'public',
                  created_at = COALESCE(created_at, %s),
                  updated_at = %s
                WHERE id = %s AND user_id = %s
                """,
                values + (now, now, row_id, user_id),
            )
        else:
            cursor.execute(
                """
                INSERT INTO risp_marketprice
                  (user_id, country, year, quarter, species, market_level, product,
                   price_min, price_max, price_avg, reference,
                   program, visibility, created_at, updated_at)
                VALUES
                  (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'public', %s, %s)
                """,
                (user_id, country, year, quarter) + values + (now, now),
            )
        imported += 1

    return imported, errors


IMPORTERS: Dict[str, Callable[..., Tuple[int, List[Dict[str, Any]]]]] = {
    "outbreaks": _import_outbreaks,
    "vaccination": _import_vaccinations,
    "marketprice": _import_market_prices,
}


async def process_upload(
    category: str,
    file: UploadFile,
    *,
    current_user: dict,
    is_soi: bool,
    get_db_connection: Callable[[], Any],
    default_year: Optional[str] = None,
    default_quarter: Optional[str] = None,
) -> Dict[str, Any]:
    if category not in IMPORTERS:
        raise HTTPException(status_code=404, detail=f"Unknown upload category: {category}")

    wb = await _read_workbook(file)
    ws = _pick_sheet(wb, category)
    user_id = current_user.get("id")
    country = current_user.get("country")
    program = "soi" if is_soi else "risp"

    connection = get_db_connection()
    cursor = connection.cursor()
    country_id = risp_templates._resolve_country_id(cursor, country)

    importer = IMPORTERS[category]
    importer_kwargs = dict(
        cursor=cursor,
        user_id=user_id,
        country=country,
        program=program,
        country_id=country_id,
        is_soi=is_soi,
        default_year=default_year,
    )
    if category != "vaccination":
        importer_kwargs["default_quarter"] = default_quarter

    imported, errors = importer(ws, **importer_kwargs)

    if errors:
        connection.rollback()
        cursor.close()
        connection.close()
        return {
            "success": False,
            "has_errors": True,
            "message": f"{len(errors)} row(s) failed validation. Nothing was imported.",
            "imported_count": 0,
            "error_count": len(errors),
            "errors": errors,
        }

    if imported == 0:
        connection.rollback()
        cursor.close()
        connection.close()
        return {
            "success": False,
            "has_errors": False,
            "message": "No data rows found in the file. Fill in at least one row and try again.",
            "imported_count": 0,
            "error_count": 0,
            "errors": [],
        }

    connection.commit()
    cursor.close()
    connection.close()

    return {
        "success": True,
        "has_errors": False,
        "message": f"Successfully imported {imported} row(s).",
        "imported_count": imported,
        "error_count": 0,
        "errors": [],
    }
